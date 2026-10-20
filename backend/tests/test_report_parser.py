import io
import json
import pytest
from PIL import Image, ImageDraw
import docx
from fastapi.testclient import TestClient

from app.main import app
from app.services.document_parser import (
    DocumentParserService,
    DocumentParserError,
    DocumentSecurityError,
)
from app.services.risk_engine import RiskEngineService


client = TestClient(app)


# -----------------------------------------------------------------------------
# 1. JSON Tests
# -----------------------------------------------------------------------------
def test_json_standard_findings_report():
    data = {
        "report_id": "REP-2024-TEST",
        "target": "api.cyberquant.internal",
        "findings": [
            {
                "id": "FIND-01",
                "title": "Remote Code Execution via Insecure Deserialization",
                "severity": "Critical",
                "cvss_score": 9.8,
                "cve": "CVE-2024-1000",
                "recommendation": "Upgrade vulnerable library to latest patch release",
            },
            {
                "id": "FIND-02",
                "title": "Missing Strict-Transport-Security Header",
                "severity": "Low",
                "cvss_score": 2.5,
                "recommendation": "Configure HSTS header on reverse proxy",
            },
        ],
    }
    content = json.dumps(data).encode("utf-8")
    resp = DocumentParserService.parse_json("findings.json", content)
    assert resp.report_id == "REP-2024-TEST"
    assert resp.target == "api.cyberquant.internal"
    assert resp.summary.critical == 1
    assert resp.summary.low == 1
    assert resp.summary.total_findings == 2
    assert len(resp.findings) == 2
    assert resp.findings[0].cve == "CVE-2024-1000"
    assert resp.findings[0].is_inferred is False
    assert "direct_loss" in resp.missing_fields  # Numbers not silently invented


def test_json_fair_risk_scenarios_calculation():
    data = {
        "target": "Customer Database Server",
        "scenarios": [
            {
                "scenario_name": "Ransomware Attack on DB",
                "asset_name": "Customer Database",
                "asset_value": 5000000.0,
                "threat_event_frequency": 2.0,
                "vulnerability_probability": 0.25,
                "direct_loss": 200000.0,
                "indirect_loss": 300000.0,
                "proposed_security_cost": 50000.0,
                "mitigated_vulnerability_probability": 0.05,
            }
        ],
    }
    content = json.dumps(data).encode("utf-8")
    resp = DocumentParserService.parse_json("scenarios.json", content)
    assert resp.detected_format == "JSON (FAIR Scenarios)"
    assert len(resp.scenarios) == 1
    sc = resp.scenarios[0]
    assert sc.metrics is not None
    # LEF = 2.0 * 0.25 = 0.5
    # SLE = 200,000 + 300,000 = 500,000
    # ALE = 0.5 * 500,000 = 250,000
    assert sc.metrics.loss_event_frequency == 0.5
    assert sc.metrics.single_loss_expectancy == 500000.0
    assert sc.metrics.annual_loss_expectancy == 250000.0
    assert sc.metrics.annual_risk_reduction is not None
    assert sc.metrics.security_roi_percentage is not None
    assert sc.metrics.security_roi_percentage > 0


def test_json_empty_and_malformed():
    with pytest.raises(DocumentParserError, match="empty"):
        DocumentParserService.parse_json("empty.json", b"   ")

    with pytest.raises(DocumentParserError, match="Malformed JSON"):
        DocumentParserService.parse_json("bad.json", b"{ invalid json")


# -----------------------------------------------------------------------------
# 2. CSV Tests
# -----------------------------------------------------------------------------
def test_csv_fair_risk_scenarios_calculation():
    csv_text = (
        "scenario_name,asset_name,asset_value,threat_event_frequency,vulnerability_probability,direct_loss,indirect_loss,proposed_security_cost,mitigated_vulnerability_probability\n"
        "SQL Injection Breach,Production Web DB,2000000,1.5,0.4,100000,150000,25000,0.1\n"
        "Credential Stuffing,Auth Portal,1000000,3.0,0.2,50000,50000,15000,0.05\n"
    )
    resp = DocumentParserService.parse_csv("scenarios.csv", csv_text.encode("utf-8"))
    assert resp.detected_format == "CSV (FAIR Risk Scenarios)"
    assert len(resp.scenarios) == 2
    sc1 = resp.scenarios[0]
    # LEF = 1.5 * 0.4 = 0.6; SLE = 250,000; ALE = 150,000
    assert sc1.metrics is not None
    assert sc1.metrics.loss_event_frequency == 0.6
    assert sc1.metrics.single_loss_expectancy == 250000.0
    assert sc1.metrics.annual_loss_expectancy == 150000.0
    assert resp.calculated_metrics is not None


def test_csv_vulnerability_findings():
    csv_text = (
        "id,title,severity,cvss_score,cve,asset,recommendation\n"
        "V-1,Open Telnet Port,High,7.5,,Edge Router,Disable Telnet service\n"
        "V-2,Heartbleed OpenSSL,Critical,9.0,CVE-2014-0160,Web Gateway,Patch OpenSSL\n"
    )
    resp = DocumentParserService.parse_csv("vulns.csv", csv_text.encode("utf-8"))
    assert resp.detected_format == "CSV (Vulnerability Findings)"
    assert resp.summary.critical == 1
    assert resp.summary.high == 1
    assert resp.summary.total_findings == 2
    assert resp.findings[1].cve == "CVE-2014-0160"


def test_csv_missing_fields_never_invented():
    # CSV with missing numerical fields: TEF and indirect_loss omitted
    csv_text = (
        "scenario_name,asset_name,asset_value,direct_loss\n"
        "Payment Loss,Payment Gateway,1000000,50000\n"
    )
    resp = DocumentParserService.parse_csv("partial.csv", csv_text.encode("utf-8"))
    assert "threat_event_frequency" in resp.missing_fields
    assert "indirect_loss" in resp.missing_fields
    assert len(resp.warnings) > 0
    # Values were NOT invented
    assert resp.scenarios[0].threat_event_frequency is None
    assert resp.scenarios[0].metrics is None


# -----------------------------------------------------------------------------
# 3. TXT Tests
# -----------------------------------------------------------------------------
def test_txt_document_parsing():
    txt_content = (
        "Penetration Test Report - Host: sec-prod-01.corp\n"
        "Asset Value: $1,500,000\n"
        "Direct Loss: $100,000\n"
        "Indirect Loss: $200,000\n"
        "Threat Event Frequency: 1.2\n"
        "Vulnerability Probability: 30%\n\n"
        "1. [CRITICAL] SQL Injection Vulnerability in Billing Module\n"
        "CVE-2023-45678 - CVSS: 9.8\n"
        "Remediation: Use parameterized SQL queries across all ORM mappings.\n"
    )
    resp = DocumentParserService.parse_txt("audit.txt", txt_content.encode("utf-8"))
    assert resp.detected_format == "TXT"
    assert resp.summary.critical >= 1
    assert len(resp.findings) >= 1
    assert resp.findings[0].cve == "CVE-2023-45678"
    assert resp.findings[0].is_inferred is False
    assert len(resp.scenarios) >= 1
    assert resp.scenarios[0].metrics is not None


# -----------------------------------------------------------------------------
# 4. DOCX Tests
# -----------------------------------------------------------------------------
def test_docx_document_parsing():
    doc = docx.Document()
    doc.add_heading("Internal Security Audit", level=1)
    doc.add_paragraph("Target: internal-auth.domain.local")
    doc.add_paragraph("[CRITICAL] Unauthenticated API Endpoint Access")
    doc.add_paragraph("CVSS: 9.1")
    doc.add_paragraph("Remediation: Enforce JWT authentication on private routes.")
    doc.add_paragraph("CVE-2024-7777 detected in middleware layer.")

    buf = io.BytesIO()
    doc.save(buf)
    content_bytes = buf.getvalue()

    resp = DocumentParserService.parse_docx("audit.docx", content_bytes)
    assert resp.detected_format == "DOCX"
    assert resp.summary.critical >= 1
    assert any(f.cve == "CVE-2024-7777" for f in resp.findings)


def test_docx_corrupt_package():
    with pytest.raises(DocumentParserError, match="Invalid DOCX|Corrupted DOCX"):
        DocumentParserService.detect_format_and_validate("corrupt.docx", b"PK\x03\x04badzipdata")


# -----------------------------------------------------------------------------
# 5. PDF Tests (Text & Scanned OCR)
# -----------------------------------------------------------------------------
@pytest.mark.anyio
async def test_pdf_embedded_text():
    content = "BT /F1 12 Tf 50 700 Td (Target: cloud-gateway.internal [HIGH] Privilege Escalation CVE-2024-3333 CVSS: 8.5) Tj ET"
    pdf_str = f"""%PDF-1.4
1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj
2 0 obj << /Type /Pages /Kids [3 0 R] /Count 1 >> endobj
3 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >> endobj
4 0 obj << /Length {len(content)} >> stream
{content}
endstream endobj
5 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj
xref
0 6
0000000000 65535 f 
0000000009 00000 n 
0000000058 00000 n 
0000000115 00000 n 
0000000244 00000 n 
0000000300 00000 n 
trailer << /Size 6 /Root 1 0 R >>
startxref
380
%%EOF"""
    resp = await DocumentParserService.parse_pdf("report.pdf", pdf_str.encode("latin-1"))
    assert resp.detected_format == "PDF (Embedded Text)"
    assert len(resp.findings) >= 1
    assert resp.ocr_status == "NOT_NEEDED"


@pytest.mark.anyio
async def test_pdf_scanned_ocr():
    # Build image-only PDF using Pillow
    from PIL import ImageFont
    font = ImageFont.load_default(size=22)
    img = Image.new("RGB", (600, 180), color=(255, 255, 255))
    d = ImageDraw.Draw(img)
    d.text((30, 40), "CVE-2024-8888 Critical Vulnerability", fill=(0, 0, 0), font=font)
    buf = io.BytesIO()
    img.save(buf, format="PDF")
    pdf_bytes = buf.getvalue()

    resp = await DocumentParserService.parse_pdf("scanned_doc.pdf", pdf_bytes)
    assert "PDF (Scanned" in resp.detected_format
    assert "USED" in resp.ocr_status
    assert len(resp.findings) >= 1


# -----------------------------------------------------------------------------
# 6. Image Tests (PNG & JPG OCR)
# -----------------------------------------------------------------------------
@pytest.mark.anyio
async def test_image_png_ocr():
    from PIL import ImageFont
    font = ImageFont.load_default(size=22)
    img = Image.new("RGB", (600, 180), color=(255, 255, 255))
    d = ImageDraw.Draw(img)
    d.text((30, 40), "CVE-2024-4444 Critical Alert", fill=(0, 0, 0), font=font)
    buf = io.BytesIO()
    img.save(buf, format="PNG")

    resp = await DocumentParserService.parse_image("scan.png", buf.getvalue(), "PNG")
    assert "Image (PNG" in resp.detected_format
    assert "USED" in resp.ocr_status
    assert len(resp.findings) >= 1


@pytest.mark.anyio
async def test_image_jpg_ocr():
    from PIL import ImageFont
    font = ImageFont.load_default(size=22)
    img = Image.new("RGB", (600, 180), color=(255, 255, 255))
    d = ImageDraw.Draw(img)
    d.text((30, 40), "CVE-2024-2222 High Risk", fill=(0, 0, 0), font=font)
    buf = io.BytesIO()
    img.save(buf, format="JPEG")

    resp = await DocumentParserService.parse_image("scan.jpg", buf.getvalue(), "JPG")
    assert "Image (JPG" in resp.detected_format
    assert "USED" in resp.ocr_status
    assert len(resp.findings) >= 1


# -----------------------------------------------------------------------------
# 7. Security Tests
# -----------------------------------------------------------------------------
def test_security_executable_signature_rejection():
    # Windows PE executable (MZ header)
    with pytest.raises(DocumentSecurityError, match="Executable binaries"):
        DocumentParserService.detect_format_and_validate("payload.exe", b"MZ\x90\x00\x03\x00\x00\x00")

    # Linux ELF executable
    with pytest.raises(DocumentSecurityError, match="Executable binaries"):
        DocumentParserService.detect_format_and_validate("malware.elf", b"\x7fELF\x02\x01\x01\x00")


def test_security_empty_file_rejection():
    with pytest.raises(DocumentParserError, match="empty"):
        DocumentParserService.detect_format_and_validate("empty.txt", b"")


# -----------------------------------------------------------------------------
# 8. FastAPI Upload Endpoint Integration Test
# -----------------------------------------------------------------------------
def test_upload_endpoint_integration():
    csv_payload = (
        "scenario_name,asset_name,asset_value,threat_event_frequency,vulnerability_probability,direct_loss,indirect_loss\n"
        "Endpoint Scenario,Cloud Storage,1000000,1.0,0.5,50000,50000\n"
    )
    files = {"file": ("test_scenario.csv", csv_payload.encode("utf-8"), "text/csv")}
    response = client.post("/api/v1/reports/upload", files=files)
    assert response.status_code == 200
    data = response.json()
    assert data["detected_format"] == "CSV (FAIR Risk Scenarios)"
    assert len(data["scenarios"]) == 1
    assert data["scenarios"][0]["metrics"]["annual_loss_expectancy"] == 50000.0


def test_upload_endpoint_unsupported_format():
    files = {"file": ("script.py", b"print('hello')", "text/x-python")}
    response = client.post("/api/v1/reports/upload", files=files)
    assert response.status_code == 400
    assert "Unsupported file format" in response.json()["detail"]
